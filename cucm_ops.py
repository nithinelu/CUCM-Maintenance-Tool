"""CUCM maintenance operations built on the platform CLI (cucm_cli) and AXL SQL.

  * cluster discovery / versions / health check
  * DRS backup to the embedded SFTP server (device add/list/delete, backup, status, history)
  * upgrades and COP file installs ("utils system upgrade initiate"), switch-version
  * data export: "file get" (TFTP, traces, logs, CDR) to SFTP, AXL SQL inventory -> Excel
"""

import os
import re
import time
from datetime import datetime

import requests
import urllib3
from lxml import etree

from cucm_cli import CLIError, Responder

HOSTKEY_YES = Responder(r"continue connecting \(yes/no(?:/\[fingerprint\])?\)\??\s*$", "yes",
                        name="accept SFTP host key")


# --------------------------------------------------------------------------
# Cluster / versions / health
# --------------------------------------------------------------------------
CLUSTER_LINE = re.compile(
    r"^\s*(\d{1,3}(?:\.\d{1,3}){3})\s+(\S+)\s+(\S+)\s+(Publisher|Subscriber)\s+(\S+)", re.M | re.I)


def discover_nodes(sh):
    """Parse 'show network cluster' on the publisher into node dicts."""
    out = sh.run("show network cluster", timeout=60)
    nodes = []
    for ip, fqdn, hostname, role, ntype in CLUSTER_LINE.findall(out):
        if any(n["host"] == ip for n in nodes):
            continue
        product = {"callmanager": "CUCM", "cups": "IM&P", "cuc": "Unity"}.get(ntype.lower(), ntype)
        nodes.append({"host": ip, "fqdn": fqdn, "hostname": hostname,
                      "role": role.capitalize(), "type": product})
    nodes.sort(key=lambda n: (n["type"] != "CUCM", n["role"] != "Publisher", n["hostname"]))
    return nodes, out


def versions(sh):
    act = sh.run("show version active", timeout=60)
    ina = sh.run("show version inactive", timeout=60)
    m1 = re.search(r"Active Master Version:\s*(\S+)", act)
    m2 = re.search(r"Inactive Master Version:\s*(\S+)", ina)
    cops = re.findall(r"^\s*(\S+\.cop\S*)\s*$", act, re.M | re.I)
    return {"active": m1.group(1) if m1 else "?", "inactive": m2.group(1) if m2 else "none",
            "cops": cops, "raw": act + "\n\n" + ina}


HEALTH_COMMANDS = [
    "show version active",
    "show version inactive",
    "show network cluster",
    "show status",
    "utils ntp status",
    "utils dbreplication runtimestate",
    "utils service list",
    "utils core active list",
    "show network eth0",
]


def health_check(sh, include_diagnose=False, on_output=None, cancel=None):
    cmds = list(HEALTH_COMMANDS) + (["utils diagnose test"] if include_diagnose else [])
    parts = []
    for cmd in cmds:
        if cancel is not None and cancel.is_set():
            break
        try:
            out = sh.run(cmd, timeout=600, on_output=on_output, cancel=cancel)
        except CLIError as e:
            out = f"ERROR: {e}"
        parts.append(f"===== {cmd} =====\n{out}\n")
    return "\n".join(parts)


# --------------------------------------------------------------------------
# DRS backup
# --------------------------------------------------------------------------
DEFAULT_FEATURES = "UCM,CDR_CAR"


def drs_device_list(sh):
    return sh.run("utils disaster_recovery device list", timeout=120)


def drs_delete_device(sh, name):
    return sh.run(f"utils disaster_recovery device delete {name}", timeout=120,
                  responders=[Responder(r"\((?:y/n|yes/no)\)\??\s*:?\s*$", "y")])


def drs_add_device(sh, name, path, sftp_ip, sftp_user, sftp_password, keep=3, replace=True):
    """Create the DRS network backup device pointing at the SFTP server."""
    resp = [HOSTKEY_YES,
            Responder(r"[Pp]assword[^\n]*:\s*$", sftp_password, secret=True, name="SFTP password")]
    cmd = f"utils disaster_recovery device add network {name} {path} {sftp_ip} {sftp_user} {int(keep)}"
    out = sh.run(cmd, responders=resp, timeout=180)
    if re.search(r"already exist", out, re.I) and replace:
        sh.log(f"[{sh.host}] backup device {name} exists - recreating it")
        drs_delete_device(sh, name)
        out = sh.run(cmd, responders=resp, timeout=180)
    if re.search(r"error|fail|invalid|unable", out, re.I) and not re.search(r"success", out, re.I):
        raise CLIError(f"DRS device add failed:\n{out}")
    return out


def drs_start_backup(sh, features, device):
    feats = ",".join(f.strip() for f in features.replace(" ", ",").split(",") if f.strip())
    out = sh.run(f"utils disaster_recovery backup network {feats} {device}", timeout=300)
    if not re.search(r"initiated|started|success", out, re.I):
        raise CLIError(f"Backup did not start:\n{out}")
    return out


def drs_status(sh, kind="backup"):
    out = sh.run(f"utils disaster_recovery status {kind}", timeout=120)
    st = {"raw": out, "status": "", "percent": None, "tar": "", "components": []}
    m = re.search(r"Status:\s*(.+)", out)
    if m:
        st["status"] = m.group(1).strip()
    m = re.search(r"Percentage Complete:\s*(\d+)", out)
    if m:
        st["percent"] = int(m.group(1))
    m = re.search(r"Tar Filename:\s*(\S+)", out)
    if m:
        st["tar"] = m.group(1)
    for line in out.splitlines():
        cols = line.split()
        if len(cols) >= 4 and cols[0].isupper() and cols[3].upper() in (
                "SUCCESS", "ERROR", "ACTIVE", "PENDING", "FAILED", "WARNING"):
            st["components"].append(cols[:4])
    return st


def drs_finished(st):
    s = st["status"].upper()
    if any(k in s for k in ("SUCCESS", "COMPLETED", "ERROR", "FAIL", "CANCEL", "WARNING")):
        return True
    return st["percent"] == 100


def drs_wait(sh, poll=15, cancel=None, on_status=None, timeout=6 * 3600):
    start = time.time()
    while True:
        if cancel is not None and cancel.is_set():
            raise CLIError("stopped watching (the backup keeps running on CUCM)")
        st = drs_status(sh)
        if on_status:
            on_status(st)
        if drs_finished(st):
            return st
        if time.time() - start > timeout:
            raise CLIError("backup still running after the time limit")
        for _ in range(poll * 10):
            if cancel is not None and cancel.is_set():
                break
            time.sleep(0.1)


def drs_history(sh):
    return sh.run("utils disaster_recovery history backup", timeout=120)


def drs_cancel(sh):
    return sh.run("utils disaster_recovery cancel_backup", timeout=120,
                  responders=[Responder(r"\((?:y/n|yes/no)\)\??\s*:?\s*$", "yes")])


def drs_schedule_list(sh):
    return sh.run("utils disaster_recovery schedule list", timeout=120)


# --------------------------------------------------------------------------
# Upgrade / COP install
# --------------------------------------------------------------------------
MENU = re.compile(r'Please select an option \(1 - \d+ or "q" \)\s*:\s*$')
MENU_ITEM = re.compile(r"^\s*(\d+)\)\s+(.+?)\s*$", re.M)


def upgrade_initiate(sh, sftp_ip, sftp_user, sftp_password, directory, filename,
                     auto_switch=False, sftp_port=22, on_output=None, cancel=None, list_only=False):
    """Install an upgrade ISO or COP file from the SFTP server on one node.

    Returns (ok, output). Answers the interactive questions of
    'utils system upgrade initiate' (10.x - 15.x wording).
    list_only: stop at the file menu (answer 'q') and return (True, [offered files]) -
    proves CUCM can log in to the SFTP server without installing anything."""
    problem = {}
    offered = []

    def pick(pending, _m):
        items = MENU_ITEM.findall(pending)
        if re.search(r"Remote Filesystem via SFTP", pending):
            for num, label in items:
                if "SFTP" in label:
                    return num
        if list_only and items:
            offered.extend(lbl for _, lbl in items if lbl.lower() != "quit")
            return "q"
        if re.search(r"Available options and upgrades|\.cop|\.iso|\.sgn", pending, re.I):
            for num, label in items:
                if label.strip() == filename or label.strip().endswith("/" + filename):
                    return num
            problem["msg"] = (f"'{filename}' is not offered by CUCM. Offered: "
                              + ", ".join(lbl for _, lbl in items if lbl.lower() != "quit"))
            return "q"
        if items and all(lbl.lower() == "quit" for _, lbl in items):
            problem["msg"] = "CUCM found no valid upgrade/COP files in that directory."
            return "q"
        return None

    resp = [
        HOSTKEY_YES,
        Responder(MENU, pick, name="menu selection"),
        Responder(r"Directory\s*:\s*$", directory),
        Responder(r"SMTP[^\n]*:\s*$", ""),
        Responder(r"(?<!SMTP )Server\s*:\s*$", sftp_ip),
        Responder(r"[Pp]ort[^\n]*:\s*$", str(sftp_port)),
        Responder(r"User ?[Nn]ame\s*:\s*$", sftp_user),
        Responder(r"[Pp]assword\s*:\s*$", sftp_password, secret=True),
    ]
    if list_only:
        # Never confirm anything in a dry run.
        resp.append(Responder(r"\((?:yes/no|y/n)\)\s*\??\s*:?\s*$", "no", name="dry run: decline"))
    else:
        resp += [
            Responder(r"[Ss]witch versions[^\n]*\(yes/no\)\s*:?\s*$", "yes" if auto_switch else "no",
                      name="auto switch version"),
            Responder(r"[Ss]tart installation \(yes/no\)\s*:?\s*$", "yes"),
            Responder(r"(?:continue|proceed)[^\n]*\((?:yes/no|y/n)\)\s*\??\s*:?\s*$", "yes"),
        ]
    sh.log(f"[{sh.host}] upgrade initiate: {directory}/{filename or '(list only)'}")
    try:
        out = sh.run("utils system upgrade initiate", responders=resp,
                     timeout=600 if list_only else None,
                     idle_timeout=180 if list_only else 3 * 3600, on_output=on_output,
                     cancel=cancel, expect_disconnect=auto_switch and not list_only)
    except CLIError as e:
        if "cancel" in str(e).lower():
            try:
                sh.run("utils system upgrade cancel", timeout=120,
                       responders=[Responder(r"\(yes/no\)\s*:?\s*$", "yes")])
            except Exception:
                pass
        raise
    if list_only:
        reached = offered or problem or re.search(r"Available options|no valid upgrade", out, re.I)
        if not reached:
            raise CLIError("CUCM did not get to the file list (SFTP login / path / firewall?)\n"
                           + out[-1200:])
        return True, offered
    if problem:
        raise CLIError(problem["msg"])
    ok = bool(re.search(r"successful|completed successfully|installed successfully|"
                        r"Upgrade completed|The upgrade .* complete", out, re.I)) and \
        not re.search(r"\b(failed|failure)\b", out, re.I)
    return ok, out


def upgrade_status(sh):
    return sh.run("utils system upgrade status", timeout=120)


def switch_version(sh, on_output=None):
    """Swap active/inactive partitions. The node restarts - session ends."""
    return sh.run("utils system switch-version",
                  responders=[Responder(r"\((?:Yes or No|yes/no)\)\s*\??\s*:?\s*$", "Yes")],
                  timeout=3600, on_output=on_output, expect_disconnect=True)


def upgrade_files(folder):
    """Upgrade ISOs / COP files available in the SFTP upgrade folder."""
    if not os.path.isdir(folder):
        return []
    out = []
    for name in sorted(os.listdir(folder)):
        p = os.path.join(folder, name)
        if os.path.isfile(p) and re.search(r"\.(iso|cop|sgn|sha512)$|\.cop\.", name, re.I):
            out.append((name, os.path.getsize(p)))
    return out


def file_hash(path, algo="sha512", progress=None):
    import hashlib
    h = hashlib.new(algo)
    total = os.path.getsize(path)
    done = 0
    with open(path, "rb") as f:
        while True:
            chunk = f.read(8 * 1024 * 1024)
            if not chunk:
                break
            h.update(chunk)
            done += len(chunk)
            if progress:
                progress(done, total)
    return h.hexdigest()


# --------------------------------------------------------------------------
# File export ("file get" -> SFTP)
# --------------------------------------------------------------------------
FILE_PRESETS = {
    "TFTP files (all)": ("tftp", "*", ""),
    "Syslog messages": ("activelog", "syslog/messages*", ""),
    "CallManager SDL traces (last 2 h)": ("activelog", "cm/trace/ccm/sdl/*", "reltime hours 2"),
    "CallManager SDI traces (last 2 h)": ("activelog", "cm/trace/ccm/sdi/*", "reltime hours 2"),
    "CTIManager SDL traces (last 2 h)": ("activelog", "cm/trace/cti/sdl/*", "reltime hours 2"),
    "CDR / CMR processed files (last 24 h)": ("activelog", "cm/cdr_repository/processed/*",
                                              "reltime hours 24 recurs"),
    "DRS (backup) logs": ("activelog", "platform/drf/trace/*", ""),
    "Install / upgrade logs": ("install", "*", ""),
    "CLI output & packet captures": ("activelog", "platform/cli/*", ""),
    "Custom": ("activelog", "", ""),
}


def file_get(sh, location, spec, options, sftp_ip, sftp_user, sftp_password, remote_dir,
             sftp_port=22, on_output=None, cancel=None):
    resp = [
        HOSTKEY_YES,
        Responder(r"Would you like to proceed \[y/n\]\??\s*$", "y"),
        Responder(r"SFTP server IP\s*:\s*$", sftp_ip),
        Responder(r"SFTP server port[^\n]*:\s*$", str(sftp_port)),
        Responder(r"User ID\s*:\s*$", sftp_user),
        Responder(r"[Pp]assword\s*:\s*$", sftp_password, secret=True),
        Responder(r"Download directory\s*:\s*$", remote_dir),
    ]
    cmd = f"file get {location} {spec} {options}".strip()
    out = sh.run(cmd, responders=resp, timeout=None, idle_timeout=1800,
                 on_output=on_output, cancel=cancel)
    if re.search(r"No files matched|not found|Invalid", out, re.I):
        raise CLIError(out.strip()[-600:])
    if re.search(r"Transfer failed|error|fail", out, re.I) and not \
            re.search(r"Transfer completed", out, re.I):
        raise CLIError(out.strip()[-600:])
    return out


# --------------------------------------------------------------------------
# Dry run: read-only checks of everything the tool needs, nothing is changed on CUCM
# --------------------------------------------------------------------------
def dry_run_node(sh, sftp_ip, sftp_user, sftp_password, sftp_port, export_dir, publisher,
                 on_output=None, cancel=None):
    """Returns [(check, ok, detail)]. Only show/status commands, the upgrade wizard up to
    its file list (answered 'q'), and one small 'file get' to this PC."""
    results = []

    def check(name, fn):
        if cancel is not None and cancel.is_set():
            return
        try:
            detail = sh.redact(str(fn()))
            results.append((name, True, detail))
        except Exception as e:
            results.append((name, False, sh.redact(str(e)).strip()[:1500]))
        sh.log(f"[{sh.host}] dry run - {name}: {'OK' if results[-1][1] else 'FAILED'}")

    results.append(("SSH login + CLI prompt", True, sh.username))

    def _ver():
        v = versions(sh)
        return f"active {v['active']}, inactive {v['inactive']}, COPs: {', '.join(v['cops']) or '-'}"
    check("Versions", _ver)
    if publisher:
        check("Cluster nodes", lambda: ", ".join(f"{n['hostname']}={n['host']} ({n['type']} {n['role']})"
                                                for n in discover_nodes(sh)[0]) or "none parsed")
        check("DRS device list (read)", lambda: drs_device_list(sh))
        check("DRS backup status (read)", lambda: drs_status(sh)["raw"])

    def _walk():
        ok, files = upgrade_initiate(sh, sftp_ip, sftp_user, sftp_password, "/upgrade", None,
                                     sftp_port=sftp_port, on_output=on_output, cancel=cancel,
                                     list_only=True)
        return "CUCM logged in to the SFTP server; offered: " + (", ".join(files) or
                                                                  "(no valid files in /upgrade)")
    check("CUCM -> SFTP: upgrade wizard file list (answered 'q')", _walk)
    check("CUCM -> SFTP: file get install system-history.log",
          lambda: file_get(sh, "install", "system-history.log", "", sftp_ip, sftp_user, sftp_password,
                           export_dir, sftp_port=sftp_port, on_output=on_output, cancel=cancel)[-400:])
    return results


# --------------------------------------------------------------------------
# AXL SQL inventory export
# --------------------------------------------------------------------------
SOAP_NS = "http://schemas.xmlsoap.org/soap/envelope/"
AXL_VERSIONS = ["15.0", "14.0", "12.5", "12.0", "11.5", "11.0", "10.5", "10.0"]

_PN = "LEFT JOIN routepartition rp ON n.fkroutepartition = rp.pkid"
SQL_PRESETS = {
    "Phones": """
        SELECT d.name, d.description, tm.name AS model, tp.name AS protocol,
               dp.name AS device_pool, css.name AS css, eu.userid AS owner_user
        FROM device d
        JOIN typeclass tc ON d.tkclass = tc.enum
        JOIN typemodel tm ON d.tkmodel = tm.enum
        JOIN typedeviceprotocol tp ON d.tkdeviceprotocol = tp.enum
        LEFT JOIN devicepool dp ON d.fkdevicepool = dp.pkid
        LEFT JOIN callingsearchspace css ON d.fkcallingsearchspace = css.pkid
        LEFT JOIN enduser eu ON d.fkenduser = eu.pkid
        WHERE tc.name = 'Phone' ORDER BY d.name""",
    "Phone lines": f"""
        SELECT d.name AS device, dnm.numplanindex AS line, n.dnorpattern AS dn,
               rp.name AS partition, dnm.display, dnm.label, dnm.e164mask
        FROM device d
        JOIN typeclass tc ON d.tkclass = tc.enum
        JOIN devicenumplanmap dnm ON dnm.fkdevice = d.pkid
        JOIN numplan n ON dnm.fknumplan = n.pkid
        {_PN}
        WHERE tc.name = 'Phone' ORDER BY d.name, dnm.numplanindex""",
    "Directory numbers": f"""
        SELECT n.dnorpattern AS dn, rp.name AS partition, n.description, n.alertingname
        FROM numplan n {_PN}
        WHERE n.tkpatternusage = 2 ORDER BY n.dnorpattern""",
    "End users": """
        SELECT userid, firstname, lastname, mailid, telephonenumber, department, manager
        FROM enduser ORDER BY userid""",
    "Route patterns": f"""
        SELECT n.dnorpattern AS pattern, rp.name AS partition, n.description,
               d.name AS destination, n.blockenable
        FROM numplan n {_PN}
        LEFT JOIN devicenumplanmap dnm ON dnm.fknumplan = n.pkid
        LEFT JOIN device d ON dnm.fkdevice = d.pkid
        WHERE n.tkpatternusage = 5 ORDER BY n.dnorpattern""",
    "Translation patterns": f"""
        SELECT n.dnorpattern AS pattern, rp.name AS partition, n.description,
               n.calledpartytransformationmask AS called_mask, n.prefixdigitsout AS prefix,
               css.name AS css
        FROM numplan n {_PN}
        LEFT JOIN callingsearchspace css ON n.fkcallingsearchspace_translation = css.pkid
        WHERE n.tkpatternusage = 3 ORDER BY n.dnorpattern""",
    "Hunt pilots": f"""
        SELECT n.dnorpattern AS pilot, rp.name AS partition, n.description
        FROM numplan n {_PN}
        WHERE n.tkpatternusage = 7 ORDER BY n.dnorpattern""",
    "Gateways & trunks": """
        SELECT d.name, d.description, tc.name AS class, tm.name AS model,
               tp.name AS protocol, dp.name AS device_pool
        FROM device d
        JOIN typeclass tc ON d.tkclass = tc.enum
        JOIN typemodel tm ON d.tkmodel = tm.enum
        JOIN typedeviceprotocol tp ON d.tkdeviceprotocol = tp.enum
        LEFT JOIN devicepool dp ON d.fkdevicepool = dp.pkid
        WHERE tc.name IN ('Gateway', 'Trunk') ORDER BY tc.name, d.name""",
    "Device pools": """
        SELECT dp.name, cmg.name AS cm_group, dts.name AS date_time_group, r.name AS region
        FROM devicepool dp
        LEFT JOIN callmanagergroup cmg ON dp.fkcallmanagergroup = cmg.pkid
        LEFT JOIN datetimesetting dts ON dp.fkdatetimesetting = dts.pkid
        LEFT JOIN region r ON dp.fkregion = r.pkid
        ORDER BY dp.name""",
    "CSS & partitions": """
        SELECT css.name AS css, css.description, m.sortorder AS position, rp.name AS partition
        FROM callingsearchspace css
        LEFT JOIN callingsearchspacemember m ON m.fkcallingsearchspace = css.pkid
        LEFT JOIN routepartition rp ON m.fkroutepartition = rp.pkid
        ORDER BY css.name, m.sortorder""",
    "Cluster servers": """
        SELECT name, description FROM processnode WHERE systemnode = 'f' ORDER BY name""",
}


class AXLSQL:
    def __init__(self, host, username, password, port=8443, timeout=120):
        self.url = f"https://{host.strip()}:{int(port)}/axl/"
        self.s = requests.Session()
        self.s.auth = (username, password)
        self.s.verify = False
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        self.timeout = timeout
        self.version = None

    def _post(self, version, sql):
        ns = f"http://www.cisco.com/AXL/API/{version}"
        body = (f'<soapenv:Envelope xmlns:soapenv="{SOAP_NS}" xmlns:ns="{ns}"><soapenv:Header/>'
                f"<soapenv:Body><ns:executeSQLQuery><sql><![CDATA[{sql}]]></sql>"
                f"</ns:executeSQLQuery></soapenv:Body></soapenv:Envelope>")
        return self.s.post(self.url, data=body.encode("utf-8"), timeout=self.timeout, headers={
            "Content-Type": "text/xml; charset=utf-8",
            "SOAPAction": f'"CUCM:DB ver={version} executeSQLQuery"'})

    def query(self, sql):
        sql = " ".join(sql.split())
        versions = [self.version] if self.version else AXL_VERSIONS
        last = ""
        for v in versions:
            try:
                r = self._post(v, sql)
            except requests.RequestException as e:
                raise CLIError(f"AXL connection failed: {e}") from e
            if r.status_code == 401:
                raise CLIError("AXL: 401 Unauthorized - check the application user and its "
                               "'Standard AXL API Access' role")
            if r.status_code == 599 or (r.status_code != 200 and "version" in r.text.lower()
                                        and "fault" not in r.text.lower()[:50]):
                last = r.text[:300]
                continue
            root = etree.fromstring(r.content)
            fault = root.find(f".//{{{SOAP_NS}}}Fault")
            if fault is not None:
                msg = fault.findtext("faultstring") or etree.tostring(fault, encoding=str)
                if "version" in msg.lower() and not self.version:
                    last = msg
                    continue
                raise CLIError(f"AXL SQL fault: {msg}")
            self.version = v
            rows = []
            for row in root.iter("row"):
                rows.append({etree.QName(c).localname: (c.text or "") for c in row})
            return rows
        raise CLIError(f"AXL: no supported schema version found ({last})")


def export_sql_to_excel(axl, names_sql, path, log=None, cancel=None):
    """names_sql: list of (sheet_name, sql). Writes one sheet per query."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Sheet", "Rows", "Exported"])
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    for name, sql in names_sql:
        if cancel is not None and cancel.is_set():
            break
        try:
            rows = axl.query(sql)
        except CLIError as e:
            if log:
                log(f"{name}: {e}")
            summary.append([name, f"ERROR: {e}", ""])
            continue
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", name)[:31])
        cols = list(rows[0].keys()) if rows else ["(no rows)"]
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = hdr_fill
            cell.alignment = Alignment(vertical="center")
            width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        summary.append([name, len(rows), datetime.now().strftime("%Y-%m-%d %H:%M")])
        if log:
            log(f"{name}: {len(rows)} row(s)")
    for c in "ABC":
        summary.column_dimensions[c].width = 26
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
    wb.save(path)
    return path
