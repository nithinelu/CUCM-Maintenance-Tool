"""Unity Connection: discovery of both clusters, per-cluster credentials, CUC health check,
DRS on the CUC publisher, and the CUPI inventory export. Uses a temp config - never the real one."""
import os
import sys
import tempfile
import threading

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fake_cucm
import cuc_ops
import cucm_ops as ops
from cucm_cli import CUCMShell

import maint_app

_tmp = tempfile.mkdtemp()
maint_app.CONFIG = os.path.join(_tmp, "config.json")
maint_app.LOG_DIR = os.path.join(_tmp, "logs")
maint_app.REPORT_DIR = os.path.join(_tmp, "reports")
maint_app.os.startfile = lambda p: print("would open:", p)
for _name in ("showinfo", "showwarning", "showerror"):      # dialogs would block the test
    setattr(maint_app.messagebox, _name, lambda *a, _n=_name, **k: print(f"[{_n}]", *a[1:2]))
maint_app.messagebox.askyesno = lambda *a, **k: True

# CUCM fake on 127.0.0.1:2231, CUC fake on 127.0.0.2:2231 - each reports itself in 'show network cluster'
fake_cucm.SFTP_PORT = 2025
fake_cucm.start(2231, "127.0.0.1", "cucm",
                "127.0.0.1 cucmpub.lab.local cucmpub Publisher callmanager DBPub authenticated\n")
fake_cucm.start(2231, "127.0.0.2", "cuc",
                "127.0.0.2 cucpub.lab.local cucpub Publisher cuc DBPub authenticated\n")
fake_cucm.start_cupi(8099)

# ---- module level: CUPI paging, single-item dicts, URI columns dropped
api = cuc_ops.CUPI("x", "cupiadmin", "cupipw", base_url="http://127.0.0.1:8099")
rows = cuc_ops.fetch_preset(api, "Users")
assert len(rows) == 7 and list(rows[0])[:3] == ["Alias", "FirstName", "LastName"] and "URI" not in rows[0], rows[0]
api_small = cuc_ops.CUPI("x", "cupiadmin", "cupipw", base_url="http://127.0.0.1:8099")
assert len(api_small.get_all("/vmrest/users", page_size=3)) == 7          # 3 pages, last has 1 item (dict)
assert cuc_ops.fetch_preset(api, "Classes of service") == [{"DisplayName": "Voice Mail User COS"}]
assert api.count("/vmrest/users") == 7 and api.version() == "14.0.1.13900-79"
try:
    cuc_ops.CUPI("x", "bad", "bad", base_url="http://127.0.0.1:8099").count("/vmrest/users")
    raise AssertionError("bad CUPI login accepted")
except ops.CLIError as e:
    assert "401" in str(e)
print("CUPI client OK")

# ---- app level
app = maint_app.App()
app.v_root.set(tempfile.mkdtemp())
app.v_port.set("2025")
app.v_ip.set("127.0.0.1")
app.v_pub.set("127.0.0.1")
app.v_osuser.set("admin")
app.v_ospw.set("adminpw")
app.v_cucpub.set("127.0.0.2")          # CUC OS admin left empty -> uses the CUCM one
app.v_cupiuser.set("cupiadmin")
app.v_cupipw.set("cupipw")
used = []


def shell(host, cluster=None):
    u, p = app.os_creds(cluster or app.cluster_of(host))
    used.append((host, cluster or app.cluster_of(host)))
    return CUCMShell(host, u, p, port=2231, log=app.log)


app.shell = shell
app.cupi = lambda: cuc_ops.CUPI("127.0.0.2", "cupiadmin", "cupipw", base_url="http://127.0.0.1:8099")
errors, done = [], []


def steps():
    try:
        app._discover()
        types = {n["host"]: (n["type"], n["cluster"], n["role"]) for n in app.nodes}
        print("nodes:", types)
        assert types == {"127.0.0.1": ("CUCM", "cucm", "Publisher"), "127.0.0.2": ("Unity", "cuc", "Publisher")}
        assert app.publisher("cuc") == "127.0.0.2" and app.publisher("cucm") == "127.0.0.1"

        # health check uses the CUC command list on the CUC node
        app.up_picker.checked = {"127.0.0.2"}
        app._health(False)
        rep = [f for f in os.listdir(maint_app.REPORT_DIR) if f.startswith("health")][0]
        text = open(os.path.join(maint_app.REPORT_DIR, rep), encoding="utf-8").read()
        assert "===== show cuc cluster status =====" in text and "Pri Active" in text
        assert "utils dbreplication runtimestate" not in text

        # DRS on the CUC publisher with the CUC feature and its own folder
        app.ui(app.v_bkcl.set, "cuc")
        app.ui(app._bk_switch)
        threading.Event().wait(0.5)
        assert app._bk_cur == "cuc" and app.v_feat.get() == "CUC" and app.v_bkdir.get() == "/backup_cuc"
        app.start_sftp_evt = None
        app._backup()
        cuc_bk = os.listdir(os.path.join(app.v_root.get(), "backup_cuc"))
        assert any(f.endswith(".tar") for f in cuc_bk), cuc_bk
        assert ("127.0.0.2", "cuc") in used

        # registered features via the Backup tab button path
        app._pub_cmd(app._drs_registration, "DRS Registered features")

        # dry run covers both clusters + CUPI
        app.up_picker.checked = {"127.0.0.1", "127.0.0.2"}
        app._dry_run()
        rep = [f for f in os.listdir(maint_app.REPORT_DIR) if f.startswith("dryrun")][0]
        text = open(os.path.join(maint_app.REPORT_DIR, rep), encoding="utf-8").read()
        print(text.split("#####")[0])
        assert "CUPI: OK (Unity Connection 14.0.1.13900-79, 7 users)" in text
        summary = text.split("#####")[0]
        assert summary.count("[OK ] Cluster nodes") == 2 and "[OK ] show cuc cluster status" in summary
        # the fake does not implement 'device list' -> must be reported as FAIL, not OK
        assert "[FAIL] DRS device list (read)" in summary

        # CUPI export through the generic Excel writer
        from openpyxl import load_workbook
        out = os.path.join(_tmp, "cuc.xlsx")
        api2 = app.cupi()
        ops.export_to_excel([("Users", "Users"), ("Custom", "/vmrest/coses")],
                            lambda a: cuc_ops.fetch_preset(api2, a), out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Summary", "Users", "Custom"] and wb["Users"].max_row == 8

        # config: CUC settings saved, passwords not saved (Remember is off), CUCM DRS kept separate
        app.ui(app._save_cfg)
        threading.Event().wait(0.5)
        import json
        cfg = json.load(open(maint_app.CONFIG, encoding="utf-8"))
        assert cfg["cuc"]["publisher"] == "127.0.0.2" and "os_password" not in cfg["cuc"]
        assert cfg["drs_dir"] == "/backup" and cfg["drs_cuc"]["dir"] == "/backup_cuc", cfg
    except Exception:
        import traceback
        errors.append(traceback.format_exc())
    done.append(1)


app.after(200, lambda: threading.Thread(target=steps, daemon=True).start())


def check():
    if done:
        app.after(500, app.quit)
    else:
        app.after(200, check)


app.after(300, check)
app.mainloop()
app.stop_sftp()
app.destroy()
if errors:
    print(errors[0])
    sys.exit(1)
print("CUC OK")
